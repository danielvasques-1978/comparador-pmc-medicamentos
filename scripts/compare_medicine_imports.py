from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path
from statistics import median

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


RATES = ["20", "18", "17", "12"]


def normalize(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", ascii_value).strip().casefold()


def comparison_key(row: dict[str, object]) -> tuple[str, str, str]:
    return (normalize(row["name"]), normalize(row["laboratory"]), normalize(row["presentation"]))


def load_rows(path: Path) -> list[dict[str, object]]:
    return json.loads(path.read_text(encoding="utf-8"))


def row_values(row: dict[str, object]) -> list[object]:
    pmc = row.get("pmc") or {}
    return [
        row.get("name", ""),
        row.get("activeIngredient", ""),
        row.get("laboratory", ""),
        row.get("kind", ""),
        row.get("presentation", ""),
        *[pmc.get(rate, "") for rate in RATES],
        row.get("sourcePage", ""),
    ]


def needs_review(row: dict[str, object], max_abs_delta: float | None = None) -> str:
    presentation = str(row.get("presentation") or "")
    reasons: list[str] = []
    if len(presentation) < 8 or re.fullmatch(r"[\d\s]+", presentation):
        reasons.append("apresentacao curta/ambigua")
    if "venda)" in presentation.casefold():
        reasons.append("texto provavelmente truncado")
    if max_abs_delta is not None and max_abs_delta > 20:
        reasons.append("variacao acima de 20%")
    return "; ".join(reasons)


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column[:200])
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 10), 48)


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit(
            "Uso: python scripts/compare_medicine_imports.py <anterior.json> <novo.json> <saida.xlsx>"
        )

    old_path = Path(sys.argv[1])
    new_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])
    old_rows = load_rows(old_path)
    new_rows = load_rows(new_path)

    old_by_key = {comparison_key(row): row for row in old_rows}
    new_by_key = {comparison_key(row): row for row in new_rows}

    added = [new_by_key[key] for key in sorted(new_by_key.keys() - old_by_key.keys())]
    removed = [old_by_key[key] for key in sorted(old_by_key.keys() - new_by_key.keys())]

    changed_rows: list[list[object]] = []
    review_rows: list[list[object]] = []
    pct_changes_by_rate: dict[str, list[float]] = {rate: [] for rate in RATES}

    for key in sorted(old_by_key.keys() & new_by_key.keys()):
        old = old_by_key[key]
        new = new_by_key[key]
        old_pmc = old.get("pmc") or {}
        new_pmc = new.get("pmc") or {}
        diffs: list[object] = []
        abs_pcts: list[float] = []
        for rate in RATES:
            old_price = old_pmc.get(rate)
            new_price = new_pmc.get(rate)
            pct = ""
            if old_price != new_price:
                if old_price and new_price:
                    pct_value = round((float(new_price) - float(old_price)) / float(old_price) * 100, 2)
                    pct = pct_value
                    abs_pcts.append(abs(pct_value))
                    pct_changes_by_rate[rate].append(pct_value)
                diffs.extend([old_price or "", new_price or "", pct])
            else:
                diffs.extend([old_price or "", new_price or "", ""])
        if any(item != "" for item in diffs[2::3]):
            max_abs_delta = max(abs_pcts) if abs_pcts else None
            review_reason = needs_review(new, max_abs_delta)
            output_row = [
                new.get("name", ""),
                new.get("activeIngredient", ""),
                new.get("laboratory", ""),
                new.get("kind", ""),
                new.get("presentation", ""),
                *diffs,
                review_reason,
            ]
            changed_rows.append(output_row)
            if review_reason:
                review_rows.append(output_row)

    headers = ["nome", "substancia", "laboratorio", "tipo", "apresentacao"]
    price_headers = []
    for rate in RATES:
        price_headers.extend([f"PMC {rate}% anterior", f"PMC {rate}% novo", f"variacao {rate}%"])

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = [
        ["fonte anterior", old_rows[0].get("source", "") if old_rows else ""],
        ["data anterior", old_rows[0].get("tableDate", "") if old_rows else ""],
        ["registros anteriores extraidos", len(old_rows)],
        ["fonte nova", new_rows[0].get("source", "") if new_rows else ""],
        ["data nova", new_rows[0].get("tableDate", "") if new_rows else ""],
        ["registros novos extraidos", len(new_rows)],
        ["entradas potenciais", len(added)],
        ["saidas potenciais", len(removed)],
        ["apresentacoes com preco alterado", len(changed_rows)],
        ["alteracoes que exigem revisao", len(review_rows)],
    ]
    for rate, values in pct_changes_by_rate.items():
        if values:
            summary.append([f"mediana variacao PMC {rate}%", median(values)])
            summary.append([f"menor variacao PMC {rate}%", min(values)])
            summary.append([f"maior variacao PMC {rate}%", max(values)])

    add_sheet(workbook, "Resumo", ["item", "valor"], summary)
    add_sheet(
        workbook,
        "Entradas potenciais",
        [*headers, *[f"PMC {rate}%" for rate in RATES], "pagina", "revisao"],
        [row_values(row) + [needs_review(row)] for row in added],
    )
    add_sheet(
        workbook,
        "Saidas potenciais",
        [*headers, *[f"PMC {rate}%" for rate in RATES], "pagina", "revisao"],
        [row_values(row) + [needs_review(row)] for row in removed],
    )
    add_sheet(workbook, "Alteracoes de preco", [*headers, *price_headers, "revisao"], changed_rows)
    add_sheet(workbook, "Revisar", [*headers, *price_headers, "revisao"], review_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Relatorio salvo em {output_path}")


if __name__ == "__main__":
    main()
